from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass
class Lot:
    lot: int
    area_m2: float
    confidence: float | None = None
    note: str = ""


@dataclass
class LotGroup:
    start: int
    end: int
    quantity: int
    unit_area: float

    @property
    def label(self) -> str:
        if self.start == self.end:
            return f"L.{self.start:02d}"
        return f"L.{self.start:02d} ao L.{self.end:02d}"


def normalize_quadra(value) -> str:
    s = str(value or "").strip()
    if not s:
        return "SEM_QUADRA"
    m = re.fullmatch(r"0*(\d{1,3})", s)
    if m:
        digits = m.group(1)
        return str(int(digits)).zfill(max(2, len(digits)))
    s = re.sub(r"^\s*QUADRA\s*", "", s, flags=re.IGNORECASE).strip()
    return s or "SEM_QUADRA"


def quadra_sort_key(q: str):
    parts = re.split(r"(\d+)", str(q))
    return tuple(int(p) if p.isdigit() else p.casefold() for p in parts)


def normalize_lots(rows: Iterable[dict | Lot]) -> list[Lot]:
    lots: list[Lot] = []
    for row in rows:
        if isinstance(row, Lot):
            item = row
        else:
            lot_raw = row.get("lote", row.get("lot"))
            area_raw = row.get("area_m2", row.get("area"))
            confidence = row.get("confidence")
            note = str(row.get("note", row.get("observacao", "")) or "")
            if lot_raw in (None, "") or area_raw in (None, ""):
                continue
            item = Lot(
                lot=int(float(str(lot_raw).replace(",", "."))),
                area_m2=float(str(area_raw).replace(".", "").replace(",", "."))
                if isinstance(area_raw, str) and "," in area_raw
                else float(area_raw),
                confidence=float(confidence) if confidence not in (None, "") else None,
                note=note,
            )
        lots.append(item)

    lots.sort(key=lambda x: x.lot)
    return lots


def validate_lots(lots: Sequence[Lot]) -> dict:
    errors: list[str] = []
    warnings: list[str] = []

    if not lots:
        errors.append("Nenhum lote foi informado.")
        return {"errors": errors, "warnings": warnings}

    numbers = [x.lot for x in lots]
    duplicates = sorted({n for n in numbers if numbers.count(n) > 1})
    if duplicates:
        errors.append("Lotes duplicados: " + ", ".join(f"L.{n:02d}" for n in duplicates))

    invalid_areas = [x.lot for x in lots if x.area_m2 <= 0]
    if invalid_areas:
        errors.append("Áreas inválidas em: " + ", ".join(f"L.{n:02d}" for n in invalid_areas))

    max_lot = max(numbers)
    expected_seq = set(range(1, max_lot + 1))
    missing = sorted(expected_seq - set(numbers))
    if missing:
        errors.append(
            "Numeração não contínua. Lotes ausentes: "
            + ", ".join(f"L.{n:02d}" for n in missing)
        )

    low_conf = [x.lot for x in lots if x.confidence is not None and x.confidence < 0.80]
    if low_conf:
        warnings.append(
            "Revise os lotes com baixa confiança: "
            + ", ".join(f"L.{n:02d}" for n in low_conf)
        )

    return {"errors": errors, "warnings": warnings}


def group_consecutive_equal_areas(lots: Sequence[Lot], decimals: int = 2) -> list[LotGroup]:
    if not lots:
        return []

    ordered = sorted(lots, key=lambda x: x.lot)
    groups: list[LotGroup] = []

    start = ordered[0].lot
    end = ordered[0].lot
    area = round(ordered[0].area_m2, decimals)

    for item in ordered[1:]:
        item_area = round(item.area_m2, decimals)
        if item.lot == end + 1 and item_area == area:
            end = item.lot
            continue

        groups.append(LotGroup(start=start, end=end, quantity=end - start + 1, unit_area=area))
        start = end = item.lot
        area = item_area

    groups.append(LotGroup(start=start, end=end, quantity=end - start + 1, unit_area=area))
    return groups


def _safe_sheet_name(quadra: str, used: set[str]) -> str:
    base = f"QUADRA {quadra}" if not str(quadra).upper().startswith("QUADRA") else str(quadra)
    base = re.sub(r"[\\/*?:\[\]]", "_", base).strip() or "QUADRA"
    base = base[:31]
    name = base
    i = 2
    while name.casefold() in {u.casefold() for u in used}:
        suffix = f"_{i}"
        name = (base[:31-len(suffix)] + suffix)
        i += 1
    used.add(name)
    return name


def _populate_quadra_block(
    ws,
    start_row: int,
    lots: Sequence[Lot],
    quadra: str,
    title: str = "ÁREAS DOS LOTES - RESIDENCIAIS",
    group_equal_consecutive: bool = True,
) -> int:
    lots = sorted(lots, key=lambda x: x.lot)
    if not lots:
        raise ValueError(f"Nenhum lote informado para a Quadra {quadra}.")

    groups = (
        group_consecutive_equal_areas(lots)
        if group_equal_consecutive
        else [LotGroup(x.lot, x.lot, 1, round(x.area_m2, 2)) for x in lots]
    )

    title_row = start_row
    header_row = start_row + 1
    first_data_row = start_row + 2
    last_data_row = first_data_row + len(groups) - 1
    total_row = last_data_row + 1

    ws.merge_cells(start_row=title_row, start_column=1, end_row=total_row, end_column=1)
    ws.cell(title_row, 1, f"QUADRA {quadra}")
    ws.merge_cells(start_row=title_row, start_column=2, end_row=title_row, end_column=5)
    ws.cell(title_row, 2, title)

    ws.cell(header_row, 2, "LOTES")
    ws.cell(header_row, 3, "QUANTIDADE (unid)")
    ws.cell(header_row, 4, "Á. UNITÁRIA (m²)")
    ws.cell(header_row, 5, "SOMATÓRIO (m²)")

    for row_idx, group in enumerate(groups, start=first_data_row):
        ws.cell(row_idx, 2, group.label)
        ws.cell(row_idx, 3, group.quantity)
        ws.cell(row_idx, 4, group.unit_area)
        ws.cell(row_idx, 5, f"=C{row_idx}*D{row_idx}")

    ws.cell(total_row, 2, "TOTAL")
    ws.cell(total_row, 3, f"=SUM(C{first_data_row}:C{last_data_row})")
    ws.cell(total_row, 4, None)
    ws.cell(total_row, 5, f"=SUM(E{first_data_row}:E{last_data_row})")

    yellow = "F8E7A0"
    gray = "D9D9D9"
    black = "000000"
    thin = Side(style="thin", color=black)

    qcell = ws.cell(title_row, 1)
    qcell.fill = PatternFill("solid", fgColor=yellow)
    qcell.font = Font(name="Arial", size=12, bold=True)
    qcell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)

    for col in range(2, 6):
        cell = ws.cell(title_row, col)
        cell.fill = PatternFill("solid", fgColor=yellow)
        cell.font = Font(name="Arial", size=15, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in range(2, 6):
        cell = ws.cell(header_row, col)
        cell.fill = PatternFill("solid", fgColor=gray)
        cell.font = Font(name="Arial", size=12, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in range(first_data_row, total_row):
        ws.cell(row, 2).font = Font(name="Arial", size=11)
        ws.cell(row, 2).alignment = Alignment(horizontal="left", vertical="center")
        for col in (3, 4, 5):
            ws.cell(row, col).font = Font(name="Arial", size=11)
            ws.cell(row, col).alignment = Alignment(horizontal="right", vertical="center")

    for col in range(2, 6):
        cell = ws.cell(total_row, col)
        cell.fill = PatternFill("solid", fgColor=gray)
        cell.font = Font(name="Arial", size=12, bold=True)
        cell.alignment = Alignment(horizontal="right" if col >= 3 else "left", vertical="center")

    for row in range(first_data_row, total_row + 1):
        ws.cell(row, 3).number_format = "0"
        ws.cell(row, 4).number_format = '#,##0.00'
        ws.cell(row, 5).number_format = '#,##0.00'

    for row in range(title_row, total_row + 1):
        for col in range(1, 6):
            ws.cell(row, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.row_dimensions[title_row].height = 24
    ws.row_dimensions[header_row].height = 22
    for row in range(first_data_row, total_row + 1):
        ws.row_dimensions[row].height = 21

    return total_row


def _configure_areas_sheet(ws, last_row: int) -> None:
    ws.sheet_view.showGridLines = False
    widths = {"A": 5.5, "B": 18.5, "C": 27.0, "D": 22.0, "E": 24.0}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "B3"
    ws.print_area = f"A1:E{last_row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.20
    ws.page_margins.right = 0.20
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25


def build_excel_bytes(
    lots: Sequence[Lot],
    quadra: str = "01",
    title: str = "ÁREAS DOS LOTES - RESIDENCIAIS",
    group_equal_consecutive: bool = True,
) -> bytes:
    lots = sorted(lots, key=lambda x: x.lot)
    if not lots:
        raise ValueError("Nenhum lote informado para gerar o Excel.")

    wb = Workbook()
    ws = wb.active
    ws.title = "ÁREAS DOS LOTES"
    q = normalize_quadra(quadra)
    total_row = _populate_quadra_block(
        ws,
        start_row=1,
        lots=lots,
        quadra=q,
        title=title,
        group_equal_consecutive=group_equal_consecutive,
    )
    _configure_areas_sheet(ws, total_row)

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def build_multi_excel_bytes(
    quadras: Mapping[str, Sequence[Lot]],
    title: str = "ÁREAS DOS LOTES - RESIDENCIAIS",
    group_equal_consecutive: bool = True,
    include_summary: bool = False,
) -> bytes:
    cleaned: dict[str, list[Lot]] = {}
    for q_raw, lots in quadras.items():
        q = normalize_quadra(q_raw)
        normalized = sorted(list(lots), key=lambda x: x.lot)
        if normalized:
            cleaned.setdefault(q, []).extend(normalized)
    if not cleaned:
        raise ValueError("Nenhuma quadra com lotes foi informada para gerar o Excel.")

    wb = Workbook()
    ws = wb.active
    ws.title = "ÁREAS DOS LOTES"

    current_row = 1
    last_total_row = 1
    ordered_quadras = sorted(cleaned, key=quadra_sort_key)
    for index, q in enumerate(ordered_quadras):
        total_row = _populate_quadra_block(
            ws,
            start_row=current_row,
            lots=cleaned[q],
            quadra=q,
            title=title,
            group_equal_consecutive=group_equal_consecutive,
        )
        last_total_row = total_row

        if index < len(ordered_quadras) - 1:
            spacer_row = total_row + 1
            ws.row_dimensions[spacer_row].height = 9
            for col in range(1, 6):
                ws.cell(spacer_row, col).fill = PatternFill(fill_type=None)
                ws.cell(spacer_row, col).border = Border()
            current_row = spacer_row + 1

    _configure_areas_sheet(ws, last_total_row)

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
